"""Import the existing Excel workbook into PostgreSQL review staging.

Usage:
  python tools/import_workbook.py --database-url postgresql://poetry:...@localhost:5432/poetry --workbook ..\古诗.xlsx
"""

import argparse
import hashlib
import json
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


def normalize_text(value: object) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--database-url")
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    workbook_bytes = args.workbook.read_bytes()
    checksum = hashlib.sha256(workbook_bytes).hexdigest()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook.active

    staged_records = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        ordinal, title, dynasty, author, body, *_ = row
        if row_number == 1 or not any(row):
            continue
        record = {
            "ordinal": ordinal,
            "title": normalize_text(title),
            "dynasty": normalize_text(dynasty),
            "author": normalize_text(author),
            "body": normalize_text(body),
        }
        errors = [field for field in ("title", "author", "body") if not record[field]]
        staged_records.append((row_number, record, errors))

    if args.dry_run:
        print(json.dumps({"rows": len(staged_records), "invalid_rows": sum(bool(errors) for _, _, errors in staged_records)}, ensure_ascii=False))
        return

    if not args.database_url:
        parser.error("--database-url is required unless --dry-run is used")

    from psycopg import connect

    with connect(args.database_url) as connection, connection.cursor() as cursor:
        cursor.execute(
            "INSERT INTO content_import_jobs (source_file, source_checksum, status) VALUES (%s, %s, 'validated') RETURNING id",
            (args.workbook.name, checksum),
        )
        job_id = cursor.fetchone()[0]
        for row_number, record, errors in staged_records:
            cursor.execute(
                """
                INSERT INTO staged_poems (import_job_id, source_row_number, raw_record, title, dynasty_name, author_name, body, validation_errors)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (job_id, row_number, json.dumps(record, ensure_ascii=False), record["title"], record["dynasty"], record["author"], record["body"], json.dumps(errors, ensure_ascii=False)),
            )
    print(json.dumps({"job_id": str(job_id), "rows": len(staged_records)}, ensure_ascii=False))


if __name__ == "__main__":
    main()